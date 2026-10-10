# -*- coding: utf-8 -*-
"""
Exports Excel (openpyxl).

Chaque fonction renvoie un dictionnaire :
  {"status": "Succès", "nom_fichier": ..., "contenu": <octets .xlsx>, ...}
ou {"status": "<message>"} quand il n'y a rien a exporter.

Le serveur encode le contenu en base64 ; le navigateur le telecharge. Il n'y a
plus de fenetre « Enregistrer sous » cote Python : c'est le navigateur du
poste qui demande l'export qui recoit le fichier.
"""

import datetime
import io

try:
    import openpyxl
    from openpyxl.chart import BarChart, Reference as ChartReference
    from openpyxl.styles import Font
    ERREUR_OPENPYXL = None
except ImportError as exc:                      # export indisponible, pas le serveur
    openpyxl = None
    ERREUR_OPENPYXL = str(exc)

#: Doit rester aligne sur CONDITIONNEMENTS dans web/js/core/constants.js
CONDITIONNEMENT_LABELS = {"carton": "carton", "boite": "boîte"}

#: Espace dont les sorties sont extraites pour la tracabilite chirurgie.
#: Doit rester aligne sur ESPACES_SUIVI_LOT dans web/js/core/constants.js
ESPACE_CHIRURGIE = "Salle de chir"


class ExportIndisponible(RuntimeError):
    pass


def _exiger_openpyxl():
    if openpyxl is None:
        raise ExportIndisponible(
            "Export Excel indisponible : le module openpyxl n'est pas installé "
            "(pip install -r requirements.txt). Détail : %s" % ERREUR_OPENPYXL)


def sanitize_excel_value(val):
    """Neutralise l'injection de formule Excel (CWE-1237) : =, +, -, @, tab, CR."""
    if isinstance(val, str) and val:
        if val.startswith(('=', '+', '-', '@', '\t', '\r')) or \
                val.lstrip().startswith(('=', '+', '-', '@')):
            return "'" + val
    return val


def describe_conditionnement(type_stockage, quantite_par_contenant, quantite):
    """'unite' -> "À l'unité" ; carton de 10, qte 32 -> "3 cartons de 10 + 2 unité(s)"."""
    label = CONDITIONNEMENT_LABELS.get(type_stockage)
    try:
        par_contenant = int(quantite_par_contenant or 1)
    except (TypeError, ValueError):
        par_contenant = 1
    if not label or par_contenant <= 1:
        return "À l'unité"
    qte = max(0, int(quantite or 0))
    contenants, vrac = divmod(qte, par_contenant)
    pluriel = "s" if contenants > 1 else ""
    if contenants and vrac:
        return f"{contenants} {label}{pluriel} de {par_contenant} + {vrac} unité(s)"
    if contenants:
        return f"{contenants} {label}{pluriel} de {par_contenant}"
    return f"{vrac} unité(s) ({label} de {par_contenant})"


def _gras(ligne):
    for cell in ligne:
        cell.font = Font(bold=True)


def _octets(wb):
    tampon = io.BytesIO()
    wb.save(tampon)
    return tampon.getvalue()


def _titre_feuille(wb, brut, defaut="Espace"):
    titre = brut
    for car in '/\\[]*?:':
        titre = titre.replace(car, "_" if car in "/\\" else "")
    titre = titre[:31] or defaut
    base, n = titre, 1
    while titre in wb.sheetnames:
        suffixe = "_%d" % n
        titre = base[:31 - len(suffixe)] + suffixe
        n += 1
    return titre


def _int(val):
    try:
        return int(float(val or 0))
    except (TypeError, ValueError):
        return 0


def _float(val):
    try:
        return float(val or 0)
    except (TypeError, ValueError):
        return 0.0


# ------------------------------------------------------------------
# Export du stock (un onglet par espace + recapitulatif)
# ------------------------------------------------------------------

def export_stock(db):
    _exiger_openpyxl()
    produits = {p.get("reference"): p for p in db.get("produits", [])}

    espaces = {}
    for s in db.get("stock", []):
        espace = s.get("utilisateur")
        if espace:
            espaces.setdefault(espace, []).append(s)

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    if not espaces:
        wb.create_sheet(title="Stock")

    recap = {}
    total_ht = total_ttc = 0.0
    entetes = ["Référence", "Nom", "Fournisseur", "Quantité", "Conditionnement",
               "Prix Unitaire HT", "Prix Unitaire TTC", "Valeur Totale HT", "Valeur Totale TTC"]

    for espace, lignes in espaces.items():
        ws = wb.create_sheet(title=_titre_feuille(wb, espace))
        ws.append(entetes)
        _gras(ws[1])
        espace_ht = espace_ttc = 0.0
        for item in lignes:
            ref = item.get("reference", "")
            p = produits.get(ref, {})
            qte = _int(item.get("quantite"))
            p_ht = _float(item.get("prix_unitaire_ht"))
            p_ttc = _float(item.get("prix_unitaire_ttc"))
            val_ht, val_ttc = p_ht * qte, p_ttc * qte
            ws.append([
                sanitize_excel_value(ref),
                sanitize_excel_value(p.get("nom", "")),
                sanitize_excel_value(item.get("fournisseur", "")),
                qte,
                sanitize_excel_value(describe_conditionnement(
                    p.get("type_stockage", "unite"), p.get("quantite_par_carton", 1), qte)),
                round(p_ht, 2), round(p_ttc, 2), round(val_ht, 2), round(val_ttc, 2),
            ])
            espace_ht += val_ht
            espace_ttc += val_ttc
        recap[espace] = (espace_ht, espace_ttc)
        total_ht += espace_ht
        total_ttc += espace_ttc

    ws = wb.create_sheet(title=_titre_feuille(wb, "Récapitulatif"))
    ws.append(["TOTAL GÉNÉRAL", round(total_ht, 2), round(total_ttc, 2)])
    _gras(ws[1])
    ws.append([])
    ws.append(["Espace", "Valeur Totale HT", "Valeur Totale TTC"])
    _gras(ws[3])
    for espace, (ht, ttc) in recap.items():
        ws.append([sanitize_excel_value(espace), round(ht, 2), round(ttc, 2)])

    return {"status": "Succès", "nom_fichier": "Export_Stock.xlsx", "contenu": _octets(wb)}


# ------------------------------------------------------------------
# Liste de courses
# ------------------------------------------------------------------

def liste_courses(db):
    """Articles en alerte, au minimum ou en dessous, et pas deja commandes."""
    _exiger_openpyxl()
    noms = {p.get("reference"): p.get("nom", "") for p in db.get("produits", [])}
    arretes = {p.get("reference") for p in db.get("produits", []) if _int(p.get("arrete"))}
    a_commander = {}
    for s in db.get("stock", []):
        if not _int(s.get("alerte_active")) or _int(s.get("en_commande")):
            continue
        if s.get("reference") in arretes:     # on ne l'achete plus
            continue
        if _int(s.get("quantite")) <= _int(s.get("stock_minimum")):
            ref = s.get("reference", "")
            if ref and ref not in a_commander:
                a_commander[ref] = (s.get("fournisseur", ""), ref, noms.get(ref, ""))

    if not a_commander:
        return {"status": "Aucun article en rupture ou sous le seuil minimum n'a été trouvé."}

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Liste de Courses"
    ws.append(["Fournisseur", "Référence", "Nom"])
    _gras(ws[1])
    for ligne in a_commander.values():
        ws.append([sanitize_excel_value(v) for v in ligne])

    return {"status": "Succès", "nom_fichier": "Liste_de_courses.xlsx",
            "contenu": _octets(wb), "references": list(a_commander.keys())}


# ------------------------------------------------------------------
# Statistiques de consommation
# ------------------------------------------------------------------

def est_consommation(type_tx, qte):
    """Doit rester aligne sur le filtre de l'apercu (web/settings.js)."""
    upper = (type_tx or "").upper()
    if "MODIFICATION" in upper or "TRANSFERT" in upper:
        return False, 0
    # Insensible a la casse : "SORTIE_STOCK", "Sortie (...)", "SORTIE"...
    if "SORTIE" in upper:
        return True, qte
    if type_tx == "AJUSTEMENT_MANUEL" and qte < 0:
        return True, abs(qte)
    return False, 0


def _jour(date_str):
    """Date (jour) d'une transaction, ou None si illisible.

    Format ISO attendu ; a defaut, on se rabat sur les 10 premiers
    caracteres (AAAA-MM-JJ) pour ne pas perdre une ligne a l'export.
    """
    texte = str(date_str or "").strip()
    try:
        return datetime.datetime.fromisoformat(texte.replace("Z", "+00:00")).date()
    except ValueError:
        pass
    try:
        return datetime.datetime.strptime(texte[:10], "%Y-%m-%d").date()
    except ValueError:
        return None


def _mois(date_str):
    jour = _jour(date_str)
    return jour.strftime("%Y-%m") if jour else None


def stats_consommation(db, references):
    _exiger_openpyxl()
    # Liste de textes attendue ; le reste est ignoré
    references = [r for r in references if isinstance(r, str)] \
        if isinstance(references, list) else []
    if not references:
        return {"status": "Erreur: Aucune référence sélectionnée."}
    choix = set(references)
    if "TOUTES" in choix:
        choix.discard("TOUTES")
        choix.update(t.get("reference", "") for t in db.get("transactions", []))

    conso = {}                                      # (ref, espace) -> {mois: qte}
    for t in db.get("transactions", []):
        ref = t.get("reference", "")
        if ref not in choix:
            continue
        ok, qte = est_consommation(t.get("type_transaction", ""), _int(t.get("quantite")))
        mois = _mois(t.get("date", "")) if ok else None
        if mois:
            cle = (ref, t.get("utilisateur", "Inconnu"))
            conso.setdefault(cle, {})
            conso[cle][mois] = conso[cle].get(mois, 0) + qte

    if not conso:
        return {"status": "Aucune donnée de consommation pour ces références."}

    noms = {p.get("reference"): p.get("nom", "") for p in db.get("produits", [])}
    tous_mois = sorted({m for d in conso.values() for m in d})

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Résumé Consommation"
    ws.append(["Référence", "Nom", "Espace", "Moyenne Mensuelle"] + tous_mois)
    _gras(ws[1])
    lignes = {}
    for (ref, espace), donnees in conso.items():
        moyenne = sum(donnees.values()) / len(tous_mois)
        ws.append([sanitize_excel_value(ref), sanitize_excel_value(noms.get(ref, "")),
                   sanitize_excel_value(espace), round(moyenne, 2)] +
                  [donnees.get(m, 0) for m in tous_mois])
        lignes[(ref, espace)] = ws.max_row

    graphiques = wb.create_sheet(title="Graphiques")
    for i, ((ref, espace), ligne) in enumerate(lignes.items()):
        chart = BarChart()
        chart.type = "col"
        chart.title = "Consommation - %s (%s)" % (noms.get(ref) or ref, espace)
        chart.style = 13
        chart.y_axis.title = "Quantité"
        chart.x_axis.title = "Mois"
        chart.height = 10
        chart.width = 18
        fin = 4 + len(tous_mois)
        chart.add_data(ChartReference(ws, min_col=5, min_row=ligne, max_col=fin, max_row=ligne),
                       titles_from_data=False, from_rows=True)
        chart.set_categories(ChartReference(ws, min_col=5, min_row=1, max_col=fin, max_row=1))
        colonne = "B" if i % 2 == 0 else "L"
        graphiques.add_chart(chart, "%s%d" % (colonne, 2 + (i // 2) * 22))

    return {"status": "Succès", "nom_fichier": "Statistiques_Consommation.xlsx",
            "contenu": _octets(wb)}


# ------------------------------------------------------------------
# Extraction chirurgie
# ------------------------------------------------------------------

def extraction_chirurgie(db, date_debut_str, date_fin_str):
    _exiger_openpyxl()
    if not date_debut_str or not date_fin_str or not isinstance(date_debut_str, str) \
            or not isinstance(date_fin_str, str):
        return {"status": "Erreur: Dates invalides."}
    try:
        debut = datetime.datetime.strptime(date_debut_str, "%Y-%m-%d").date()
        fin = datetime.datetime.strptime(date_fin_str, "%Y-%m-%d").date()
    except ValueError:
        return {"status": "Erreur: Format de date invalide."}

    noms = {p.get("reference"): p.get("nom", "") for p in db.get("produits", [])}
    sorties = []
    for t in db.get("transactions", []):
        if t.get("utilisateur", "") != ESPACE_CHIRURGIE:
            continue
        qte = _int(t.get("quantite"))
        type_tx = (t.get("type_transaction", "") or "").upper()
        # Un ajustement manuel n'est une sortie que s'il est negatif.
        est_sortie = ((type_tx.startswith("SORTIE") and "MODIFICATION" not in type_tx
                       and "TRANSFERT" not in type_tx)
                      or (type_tx == "AJUSTEMENT_MANUEL" and qte < 0))
        if not est_sortie:
            continue
        jour = _jour(t.get("date", ""))
        if jour and debut <= jour <= fin:
            sorties.append((jour, t, abs(qte)))

    if not sorties:
        return {"status": "Aucun article sorti de la chirurgie sur cette période."}

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Extraction Chirurgie"
    ws.append(["Date", "Référence", "Nom", "Quantité", "Type de transaction",
               "Numéro de lot", "Date de péremption"])
    _gras(ws[1])
    for jour, t, qte in sorties:
        ref = t.get("reference", "")
        ws.append([sanitize_excel_value(jour.strftime("%d/%m/%Y")),
                   sanitize_excel_value(ref), sanitize_excel_value(noms.get(ref, "")),
                   qte, sanitize_excel_value(t.get("type_transaction", "")),
                   sanitize_excel_value(t.get("lot", "")),
                   sanitize_excel_value(t.get("peremption_sortie", ""))])

    return {"status": "Succès",
            "nom_fichier": "Extraction_Chirurgie_%s_au_%s.xlsx" % (date_debut_str, date_fin_str),
            "contenu": _octets(wb)}
