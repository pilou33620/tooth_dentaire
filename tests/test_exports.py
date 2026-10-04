# -*- coding: utf-8 -*-
"""Tests des exports Excel (python/exports.py)."""

import io

import openpyxl
import pytest

import exports


def classeur(resultat):
    assert resultat["status"] == "Succès"
    return openpyxl.load_workbook(io.BytesIO(resultat["contenu"]))


def db_exemple():
    return {
        "produits": [
            {"reference": "A", "nom": "Compresses", "type_stockage": "carton", "quantite_par_carton": 10},
            {"reference": "B", "nom": "=HYPERLINK(\"x\")", "type_stockage": "unite", "quantite_par_carton": 1},
        ],
        "stock": [
            {"reference": "A", "utilisateur": "Commun", "quantite": 32, "prix_unitaire_ht": 1.0,
             "prix_unitaire_ttc": 1.2, "fournisseur": "F1", "alerte_active": 1,
             "stock_minimum": 40, "en_commande": 0},
            {"reference": "B", "utilisateur": "Cabinet 1", "quantite": 2, "prix_unitaire_ht": 5,
             "prix_unitaire_ttc": 6, "fournisseur": "F2", "alerte_active": 1,
             "stock_minimum": 1, "en_commande": 0},
            {"reference": "A", "utilisateur": "Salle de chir", "quantite": 0, "prix_unitaire_ht": 1,
             "prix_unitaire_ttc": 1.2, "fournisseur": "F1", "alerte_active": 1,
             "stock_minimum": 1, "en_commande": 1},
        ],
        "transactions": [
            {"date": "2026-09-01T10:00:00Z", "reference": "A", "utilisateur": "Commun",
             "type_transaction": "SORTIE_STOCK", "quantite": 3},
            {"date": "2026-09-15T10:00:00Z", "reference": "A", "utilisateur": "Commun",
             "type_transaction": "AJUSTEMENT_MANUEL", "quantite": -2},
            {"date": "2026-10-01T10:00:00Z", "reference": "A", "utilisateur": "Commun",
             "type_transaction": "Sortie (Transfert)", "quantite": 5},
            {"date": "2026-10-02T10:00:00Z", "reference": "A", "utilisateur": "Salle de chir",
             "type_transaction": "SORTIE_STOCK", "quantite": 1, "lot": "L1",
             "peremption_sortie": "01/01/2030"},
            {"date": "2026-10-02T11:00:00Z", "reference": "A", "utilisateur": "Salle de chir",
             "type_transaction": "AJUSTEMENT_MANUEL", "quantite": 4},
        ],
    }


@pytest.mark.parametrize("valeur, attendu", [
    ("=SUM(A1)", "'=SUM(A1)"), ("+33", "'+33"), ("-1", "'-1"), ("@x", "'@x"),
    ("  =x", "'  =x"), ("normal", "normal"), ("", ""), (12, 12)])
def test_neutralisation_des_formules(valeur, attendu):
    assert exports.sanitize_excel_value(valeur) == attendu


@pytest.mark.parametrize("type_, par, qte, attendu", [
    ("unite", 1, 5, "À l'unité"),
    ("carton", 10, 32, "3 cartons de 10 + 2 unité(s)"),
    ("carton", 10, 10, "1 carton de 10"),
    ("boite", 12, 5, "5 unité(s) (boîte de 12)"),
    ("carton", 1, 5, "À l'unité"),
    ("carton", "x", 5, "À l'unité"),
])
def test_conditionnement_lisible(type_, par, qte, attendu):
    assert exports.describe_conditionnement(type_, par, qte) == attendu


def test_export_stock_un_onglet_par_espace_et_recap():
    wb = classeur(exports.export_stock(db_exemple()))
    assert wb.sheetnames == ["Commun", "Cabinet 1", "Salle de chir", "Récapitulatif"]
    ws = wb["Commun"]
    assert ws["E2"].value == "3 cartons de 10 + 2 unité(s)"
    assert ws["H2"].value == 32.0
    recap = wb["Récapitulatif"]
    assert recap["A1"].value == "TOTAL GÉNÉRAL"
    assert recap["B1"].value == pytest.approx(42.0)


def test_export_stock_neutralise_les_noms_dangereux():
    wb = classeur(exports.export_stock(db_exemple()))
    assert wb["Cabinet 1"]["B2"].value.startswith("'=")


def test_export_stock_vide():
    wb = classeur(exports.export_stock({"stock": [], "produits": []}))
    assert "Stock" in wb.sheetnames


def test_liste_de_courses():
    resultat = exports.liste_courses(db_exemple())
    wb = classeur(resultat)
    # A en Commun est sous le seuil ; la ligne en commande est ignoree ;
    # B (2 > 1) n'est pas a commander.
    assert resultat["references"] == ["A"]
    assert wb.active["C2"].value == "Compresses"


def test_liste_de_courses_ignore_les_produits_arretes():
    db = db_exemple()
    db["produits"][0]["arrete"] = 1
    assert "Aucun article" in exports.liste_courses(db)["status"]


def test_liste_de_courses_vide():
    db = db_exemple()
    for s in db["stock"]:
        s["alerte_active"] = 0
    assert "Aucun article" in exports.liste_courses(db)["status"]


def test_stats_consommation():
    wb = classeur(exports.stats_consommation(db_exemple(), ["TOUTES"]))
    ws = wb["Résumé Consommation"]
    entetes = [c.value for c in ws[1]]
    assert entetes[:4] == ["Référence", "Nom", "Espace", "Moyenne Mensuelle"]
    lignes = {(r[0].value, r[2].value): [c.value for c in r] for r in ws.iter_rows(min_row=2)}
    # Commun : 3 en septembre + 2 (ajustement negatif) ; le transfert est exclu
    commun = lignes[("A", "Commun")]
    assert commun[entetes.index("2026-09")] == 5
    assert commun[entetes.index("2026-10")] == 0
    assert len(wb["Graphiques"]._charts) == len(lignes)


def test_stats_sans_reference():
    assert exports.stats_consommation(db_exemple(), [])["status"].startswith("Erreur")


def test_stats_sans_donnee():
    assert "Aucune donnée" in exports.stats_consommation(db_exemple(), ["Z"])["status"]


@pytest.mark.parametrize("type_tx, qte, attendu", [
    ("SORTIE_STOCK", 3, (True, 3)),
    ("Sortie (Transfert)", 3, (False, 0)),
    ("SORTIE (Modification Réf/Espace)", 3, (False, 0)),
    ("AJUSTEMENT_MANUEL", -2, (True, 2)),
    ("AJUSTEMENT_MANUEL", 2, (False, 0)),
    ("ENTREE_FACTURE", 10, (False, 0)),
])
def test_filtre_de_consommation(type_tx, qte, attendu):
    assert exports.est_consommation(type_tx, qte) == attendu


def test_extraction_chirurgie():
    resultat = exports.extraction_chirurgie(db_exemple(), "2026-10-01", "2026-10-31")
    ws = classeur(resultat).active
    lignes = list(ws.iter_rows(min_row=2, values_only=True))
    # Seule la sortie compte : l'ajustement positif est une entree
    assert lignes == [("02/10/2026", "A", "Compresses", 1, "SORTIE_STOCK", "L1", "01/01/2030")]
    assert resultat["nom_fichier"] == "Extraction_Chirurgie_2026-10-01_au_2026-10-31.xlsx"


def test_extraction_chirurgie_periode_vide():
    assert "Aucun article" in exports.extraction_chirurgie(
        db_exemple(), "2020-01-01", "2020-01-31")["status"]


@pytest.mark.parametrize("debut, fin", [("", "2026-01-01"), ("01/01/2026", "2026-01-02")])
def test_extraction_chirurgie_dates_invalides(debut, fin):
    assert exports.extraction_chirurgie(db_exemple(), debut, fin)["status"].startswith("Erreur")


def test_export_indisponible_sans_openpyxl(monkeypatch):
    monkeypatch.setattr(exports, "openpyxl", None)
    with pytest.raises(exports.ExportIndisponible):
        exports.export_stock({})
